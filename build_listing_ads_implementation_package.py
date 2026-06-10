from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"E:\亚马逊Codex独立团")
OUT = ROOT / "outputs" / "b0f2_ad_sales_analysis" / "listing_ads_implementation"
SEARCH_PATH = Path(r"C:\Users\admin\Desktop\商品推广_搜索词_报告.xlsx")
PROMOTED_PATH = Path(r"C:\Users\admin\Desktop\商品推广_推广的商品_报告.xlsx")
BACKEND_PATH = Path(r"C:\Users\admin\Downloads\商品分析V2_多站点_20260507-20260605.xlsx")
COPY_PATH = Path(r"C:\Users\admin\Desktop\capybara文案.txt")

TARGET_ASINS = ["B0DHJK7NNF", "B0F2LQNTDF"]
COMPETITOR_ASINS = ["B0FJ5VT79B", "B0FYPGN85T", "B0DG96272J", "B0FSRQ698D"]


def money(x) -> str:
    return f"${float(x):,.2f}" if pd.notna(x) else ""


def pct(x) -> str:
    return f"{float(x):.1%}" if pd.notna(x) else ""


def read_ad(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path)
    for col in ["展示量", "点击量", "花费", "7天总销售额", "7天总订单数(#)", "7天总销售量(#)"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
    return df


def classify(row: pd.Series) -> tuple[str, str, str, str]:
    term = str(row["客户搜索词"]).lower()
    clicks = float(row["点击量"])
    spend = float(row["花费"])
    sales = float(row["销售额"])
    orders = float(row["订单"])
    acos = spend / sales if sales > 0 else None

    competitor_noise = ["pokemon", "pusheen", "jellycat", "ty beanie", "sylveon"]
    if any(x in term for x in competitor_noise):
        return "否定精准", "品牌/竞品相关性弱", "立刻否定", "避免无效点击"
    if orders == 0 and (clicks >= 8 or spend >= 5.81):
        return "否定精准", "点击或花费已超过测试阈值仍无单", "立刻否定或暂停", "止损"
    if orders >= 2 and acos is not None and acos <= 0.25:
        return "精准放量", "ACOS≤25%且订单≥2", "加价10%-20%，独立精准活动承接", "放量"
    if orders >= 1 and acos is not None and acos <= 0.40:
        return "保留观察", "有转化且ACOS可接受或接近目标", "保留，7天后复盘", "稳投"
    if orders >= 1 and acos is not None and acos > 0.40:
        return "降价控量", "有单但ACOS>40%", "降价20%-35%，拆出窄词测试", "控亏"
    if clicks >= 3:
        return "低价测试", "数据不足但有点击", "降至低竞价继续观察", "小额测试"
    return "观察", "数据不足", "不主动加价", "观察"


def write_df(ws, df: pd.DataFrame, start_row=1, start_col=1):
    for j, col in enumerate(df.columns, start_col):
        c = ws.cell(start_row, j, col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, row) in enumerate(df.iterrows(), start_row + 1):
        for j, col in enumerate(df.columns, start_col):
            val = row[col]
            if pd.isna(val):
                val = ""
            ws.cell(i, j, val)


def style(ws):
    ws.freeze_panes = "A2"
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(c.value or "")) for c in col[:120]) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 42)
    ws.auto_filter.ref = ws.dimensions


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    search = read_ad(SEARCH_PATH)
    promoted = read_ad(PROMOTED_PATH)
    backend = pd.read_excel(BACKEND_PATH)
    current_copy = COPY_PATH.read_text(encoding="utf-8-sig").strip()

    for col in ["销售额", "广告花费", "毛利润", "毛利率", "ACoAS", "ACoS", "可售天数", "广告订单占比", "访问转化率", "自然点击量"]:
        if col in backend.columns:
            backend[col] = pd.to_numeric(backend[col], errors="coerce")

    campaigns = promoted[promoted["广告ASIN"].astype(str).str.upper().isin(TARGET_ASINS)]["广告活动名称"].dropna().unique().tolist()
    sub = search[search["广告活动名称"].isin(campaigns)].copy()
    kw = (
        sub.groupby("客户搜索词", dropna=False)
        .agg(展示量=("展示量", "sum"), 点击量=("点击量", "sum"), 花费=("花费", "sum"), 销售额=("7天总销售额", "sum"), 订单=("7天总订单数(#)", "sum"))
        .reset_index()
    )
    kw["CTR"] = kw["点击量"] / kw["展示量"].replace(0, pd.NA)
    kw["CPC"] = kw["花费"] / kw["点击量"].replace(0, pd.NA)
    kw["CVR"] = kw["订单"] / kw["点击量"].replace(0, pd.NA)
    kw["ACOS"] = kw["花费"] / kw["销售额"].replace(0, pd.NA)
    actions = kw.apply(classify, axis=1)
    kw["动作"] = [x[0] for x in actions]
    kw["判定依据"] = [x[1] for x in actions]
    kw["执行方式"] = [x[2] for x in actions]
    kw["目标"] = [x[3] for x in actions]
    kw = kw.sort_values(["动作", "花费"], ascending=[True, False])

    campaign_plan = pd.DataFrame(
        [
            ["SP-Exact-Profitable-DressUp", "精准盈利词组", "dress up capybara; capybara gifts for girls; stuffed capybara; gifts for girls; pink capybara plush", "35%", "按当前可承受CPC起投；出单词+10%-20%", "主承接高转化长尾"],
            ["SP-Exact-Core-Control", "核心大词控量组", "capybara plush; capybara; capybara stuffed animal", "25%", "现有竞价下调15%-25%", "排名防守，不承担利润目标"],
            ["SP-Phrase-Scenario-Longtail", "场景长尾组", "cute plushies for girls; toys for girls 8-10; cute toys for girls 10-12; birthday gifts for girls; classroom rewards", "20%", "低竞价测试，7天复盘", "礼物/课堂/年龄段场景扩词"],
            ["SP-Product-Competitor", "商品投放组", "; ".join(COMPETITOR_ASINS), "20%", "低于关键词广告竞价20%-30%", "拦截竞品详情页流量"],
        ],
        columns=["建议活动名", "结构", "投放对象", "预算占比", "竞价规则", "目的"],
    )

    negative = kw[kw["动作"].eq("否定精准")][["客户搜索词", "展示量", "点击量", "花费", "销售额", "订单", "判定依据", "执行方式"]].sort_values("花费", ascending=False)
    scale = kw[kw["动作"].isin(["精准放量", "保留观察"])][["客户搜索词", "展示量", "点击量", "花费", "销售额", "订单", "ACOS", "CVR", "动作", "执行方式"]].sort_values(["动作", "订单"], ascending=[True, False])

    listing = pd.DataFrame(
        [
            ["Title", "Capybara Plush with 9 Clothes & Accessories, 10.2\" Dress Up Stuffed Animal Toy, Cute Capybara Plushie Gift for Girls Boys Kids, Birthday Classroom Rewards"],
            ["Bullet 1", "10-Piece Dress-Up Set: Includes 1 soft capybara plush and 9 mini clothes/accessories for mix-and-match styling, so kids can create new looks again and again."],
            ["Bullet 2", "Creative Pretend Play: Children can dress, restyle, and invent little stories with their capybara, helping encourage imagination and hands-on play."],
            ["Bullet 3", "Soft & Huggable: Made with soft plush fabric and fluffy filling, this 10.2-inch capybara stuffed animal is easy to cuddle, display, and carry."],
            ["Bullet 4", "Gift for Kids & Capybara Fans: A cute choice for birthdays, classroom rewards, party prizes, holidays, and everyday surprises for boys and girls."],
            ["Bullet 5", "Easy Outfit Matching: Hats, sweaters, scarf, and mini bags make this more interactive than a regular stuffed animal and fun for capybara lovers."],
            ["Backend Search Terms", "capybara plush stuffed animal dress up toy with clothes cute plushies for girls birthday gifts classroom rewards kawaii stuffy capibara peluche"],
        ],
        columns=["模块", "优化成稿"],
    )

    image_brief = pd.DataFrame(
        [
            ["主图", "白底主体+9件配件", "保留白底，重排为水豚主体+9件配件完整铺开，主体占画面75%-85%", "让买家第一眼看懂可换装套装", "最高"],
            ["图2", "10 PCS SET", "标题改为：10 PCS SET: 1 Plush + 9 Clothes & Accessories；配件逐件编号", "强化套装价值，降低价格阻力", "最高"],
            ["图3", "Mix & Match Dress-Up Play", "用3-4套穿搭组合替代单一生活场景，展示换装玩法", "承接dress up capybara高转化词", "高"],
            ["图4", "Creative Pretend Play", "原文案改为 Inspire Imagination & Creativity 或 Creative Pretend Play", "修复英文表达，提高专业度", "高"],
            ["图5", "Gift场景", "标题改为 Birthday Gifts / Classroom Rewards / Holiday Surprises", "承接gifts for girls和classroom rewards", "高"],
            ["图6", "细节图", "Soft Fabric / Fluffy Filling / Fine Stitching；替换FULL-FIL、FINE_CRAFT", "修复不自然英文，突出材质工艺", "最高"],
            ["图7", "What You Get", "标题改为 What You Get；减少橙色边框压迫，画面更干净", "清晰展示配件清单", "中"],
            ["图8", "礼品/季节", "弱化圣诞单一场景，保留为Holiday/Birthday gift通用图", "避免非旺季相关性下降", "中"],
        ],
        columns=["图片", "新主题", "执行说明", "目的", "优先级"],
    )

    review_rules = pd.DataFrame(
        [
            ["点击≥8且无单", "否定精准", "每7天执行一次", "停止无效点击"],
            ["ACOS≤25%且订单≥2", "加价10%-20%，加入精准活动", "每7天执行一次", "放大利润词"],
            ["ACOS 25%-40%", "保留观察或微降5%-10%", "每7天执行一次", "稳定订单"],
            ["ACOS>40%", "降价20%-35%或拆词重投", "每7天执行一次", "控亏"],
            ["库存<14天", "不激进放量，只保留高转化词", "每日检查", "防断货"],
        ],
        columns=["触发条件", "动作", "频率", "目的"],
    )

    asin_budget = backend[backend["ASIN"].astype(str).str.upper().isin(["B0DHJK7NNF", "B0F2LQNTDF", "B0F2MGP1WN", "B0F2MBQVNN", "B0F2MGST2F", "B0F2M53ZX4"])][
        ["ASIN", "销售额", "广告花费", "毛利润", "毛利率", "ACoAS", "ACoS", "可售天数", "广告订单占比", "访问转化率", "自然点击量"]
    ].copy()
    asin_budget["预算动作"] = asin_budget["ASIN"].map(
        {
            "B0DHJK7NNF": "砍50%-80%，只留盈利精准词",
            "B0F2LQNTDF": "库存<21天前保守控量",
            "B0F2MGP1WN": "承接部分预算，低CPC精准放量",
            "B0F2MBQVNN": "低CPC长尾测试",
            "B0F2MGST2F": "低CPC长尾测试",
            "B0F2M53ZX4": "低CPC长尾测试",
        }
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "执行总览"
    ws["A1"] = "Capybara Listing & Ads Implementation Package"
    ws["A1"].font = Font(bold=True, size=18, color="1F4E78")
    ws["A2"] = "目标：止损B0DHJK7NNF，保护B0F2LQNTDF利润与库存，重建高意图词+高转化图片文案+库存约束的投放结构。"
    ws.merge_cells("A1:H1")
    ws.merge_cells("A2:H2")
    ws["A2"].alignment = Alignment(wrap_text=True)
    overview = pd.DataFrame(
        [
            ["第一优先级", "B0DHJK7NNF广告止损", "广告花费高、毛利润为负，先砍50%-80%预算"],
            ["第二优先级", "图片英文与主题修复", "修复语法，前置10件套、换装、礼品、课堂奖励"],
            ["第三优先级", "新建4组广告结构", "精准盈利词、核心大词控量、场景长尾、竞品商品投放"],
            ["第四优先级", "7天复盘", "按点击、订单、ACOS和库存执行加价/降价/否定"],
        ],
        columns=["优先级", "模块", "动作"],
    )
    write_df(ws, overview, 4, 1)
    style(ws)

    sheets = [
        ("广告活动结构", campaign_plan),
        ("搜索词执行动作", kw),
        ("否定词清单", negative),
        ("放量&保留词", scale),
        ("ASIN预算迁移", asin_budget),
        ("Listing文案成稿", listing),
        ("图片优化Brief", image_brief),
        ("7天复盘规则", review_rules),
    ]
    for title, df in sheets:
        w = wb.create_sheet(title)
        write_df(w, df)
        style(w)

    xlsx = OUT / "Capybara_广告调整与产品优化落地包.xlsx"
    wb.save(xlsx)

    md = OUT / "Capybara_广告调整与产品优化执行手册.md"
    md.write_text(
        f"""# Capybara 广告调整与产品优化执行手册

## 1. 当前问题
- `B0DHJK7NNF` 是第一止损对象：广告花费高、毛利润为负，应先砍 `50%-80%` 预算。
- `B0F2LQNTDF` 有利润和自然点击，但库存低，库存恢复到 `21天+` 前不激进放量。
- 大词 `capybara plush / capybara / capybara squishy / capybara gifts` 有量但ACOS偏高。
- 窄词 `dress up capybara / capybara gifts for girls / stuffed capybara / gifts for girls` 更适合作为Listing和广告主线。

## 2. 广告立即执行
1. 对 `capybara plush` 降竞价 `15%-25%`，只做排名防守。
2. 对 `capybara squishy` 降竞价 `20%-35%`。
3. 对 `capybara gifts` 降竞价 `15%-25%`，改用更窄礼品词承接。
4. 将无单或弱相关词加入否定精准：`capybara toys`、`capybara with clothes`、`capybara clothes`、`capybara things`、`pokemon plushies`、`pusheen plush`、`jellycat capybara`。
5. 新建4组活动：精准盈利词组、核心大词控量组、场景长尾组、竞品商品投放组。

## 3. Listing成稿
**Title**
Capybara Plush with 9 Clothes & Accessories, 10.2" Dress Up Stuffed Animal Toy, Cute Capybara Plushie Gift for Girls Boys Kids, Birthday Classroom Rewards

**Bullets**
1. 10-Piece Dress-Up Set: Includes 1 soft capybara plush and 9 mini clothes/accessories for mix-and-match styling, so kids can create new looks again and again.
2. Creative Pretend Play: Children can dress, restyle, and invent little stories with their capybara, helping encourage imagination and hands-on play.
3. Soft & Huggable: Made with soft plush fabric and fluffy filling, this 10.2-inch capybara stuffed animal is easy to cuddle, display, and carry.
4. Gift for Kids & Capybara Fans: A cute choice for birthdays, classroom rewards, party prizes, holidays, and everyday surprises for boys and girls.
5. Easy Outfit Matching: Hats, sweaters, scarf, and mini bags make this more interactive than a regular stuffed animal and fun for capybara lovers.

**Backend Search Terms**
capybara plush stuffed animal dress up toy with clothes cute plushies for girls birthday gifts classroom rewards kawaii stuffy capibara peluche

## 4. 图片改图顺序
1. 主图：白底，水豚主体 + 9件配件完整铺开，主体占画面75%-85%。
2. 图2：`10 PCS SET: 1 Plush + 9 Clothes & Accessories`。
3. 图3：`Mix & Match Dress-Up Play`，展示3-4套穿搭。
4. 图4：`Creative Pretend Play` 或 `Inspire Imagination & Creativity`。
5. 图5：`Birthday Gifts / Classroom Rewards / Holiday Surprises`。
6. 图6：`Soft Fabric / Fluffy Filling / Fine Stitching`。
7. 图7：`What You Get`，配件清单更干净。
8. 图8：从圣诞单一场景改成通用Holiday/Birthday gift。

## 5. 7天复盘规则
- 点击≥8且无单：否定精准。
- ACOS≤25%且订单≥2：加价10%-20%。
- ACOS 25%-40%：保留观察或微降5%-10%。
- ACOS>40%：降价20%-35%或拆词重投。
- 库存<14天：只保留高转化词，不放量。

## 6. 文件说明
- 工作簿包含：广告活动结构、搜索词执行动作、否定词清单、放量&保留词、ASIN预算迁移、Listing文案成稿、图片优化Brief、7天复盘规则。

## 7. 当前原文备份
```text
{current_copy}
```
""",
        encoding="utf-8-sig",
    )

    for name, df in [
        ("搜索词执行动作.csv", kw),
        ("否定词清单.csv", negative),
        ("广告活动结构.csv", campaign_plan),
        ("图片优化Brief.csv", image_brief),
        ("Listing文案成稿.csv", listing),
    ]:
        df.to_csv(OUT / name, index=False, encoding="utf-8-sig")

    print(f"XLSX={xlsx}")
    print(f"MD={md}")
    print(f"search_terms={len(kw)}")
    print(f"negative_terms={len(negative)}")
    print(f"scale_terms={len(scale)}")


if __name__ == "__main__":
    main()
