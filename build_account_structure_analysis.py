from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_DIR = Path(r"E:\亚马逊Codex独立团\outputs\b0f2_ad_sales_analysis")
BACKEND_PATH = Path(r"C:\Users\admin\Downloads\商品分析V2_多站点_20260507-20260605.xlsx")
PROMOTED_PATH = Path(r"C:\Users\admin\Desktop\商品推广_推广的商品_报告.xlsx")
SEARCH_PATH = Path(r"C:\Users\admin\Desktop\商品推广_搜索词_报告.xlsx")
TARGET_PARENT = "B0FXRZ3YMT"
TARGET_ASIN = "B0F2LQNTDF"


TEXT_COLS = {
    "ASIN",
    "SKU",
    "店铺",
    "父ASIN",
    "产品标题",
    "图片链接",
    "币种",
    "大类目名称",
    "子类目名称",
    "第二子类目名称",
}


def read_backend() -> pd.DataFrame:
    df = pd.read_excel(BACKEND_PATH)
    for col in df.columns:
        if col not in TEXT_COLS:
            converted = pd.to_numeric(df[col], errors="coerce")
            if converted.notna().any():
                df[col] = converted
    return df


def read_ad(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path)
    for col in ["展示量", "点击量", "花费", "7天总销售额", "7天总订单数(#)", "7天总销售量(#)"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
    return df


def pct(x) -> str:
    return "" if pd.isna(x) else f"{float(x):.1%}"


def money(x) -> str:
    return "" if pd.isna(x) else f"${float(x):,.2f}"


def account_role(row: pd.Series) -> str:
    sales = row.get("销售额", 0) or 0
    spend = row.get("广告花费", 0) or 0
    profit = row.get("毛利润", 0) or 0
    margin = row.get("毛利率", 0)
    acoas = row.get("ACoAS", 0)
    stock_days = row.get("可售天数", 0)
    conversion = row.get("访问转化率", 0)

    if stock_days < 14 and sales > 300:
        return "控量防断货"
    if profit < 0 or (spend > 30 and (pd.isna(margin) or margin < 0.08)):
        return "止损收缩"
    if stock_days > 180 and sales < 500:
        return "清库存/低成本激活"
    if sales > 300 and margin >= 0.25 and acoas <= 0.08 and stock_days >= 21:
        return "利润型放量"
    if conversion < 0.03 and sales < 300:
        return "Listing承接修复"
    return "稳态观察"


def budget_action(row: pd.Series) -> tuple[str, str, str]:
    role = row["结构角色"]
    asin = row["ASIN"]
    if role == "控量防断货":
        return "-20%~-40%", "库存恢复前限制预算，保留高转化精准词", "先补货，避免断货掉排名"
    if role == "止损收缩":
        return "-50%~-80%", "砍高ACOS活动，只留品牌/高意图精准词", "减少父体利润被广告吃掉"
    if role == "利润型放量":
        return "+20%~+40%", "增加精准词和商品投放预算", "承接自然点击，扩大高毛利销售"
    if role == "清库存/低成本激活":
        return "+0%~+15%", "只做低CPC长尾词、coupon或轻促销", "清库存，不为低转化高价买量"
    if role == "Listing承接修复":
        return "0%或暂停", "先修主图/价格/卖点，再恢复测试", "避免把流量浪费在低转化页面"
    if asin == TARGET_ASIN:
        return "-10%~+10%", "库存紧，保留盈利词，小幅优化", "利润可接受但广告ACOS偏高"
    return "维持", "继续观察7天", "数据未显示强动作需求"


def write_df(ws, df: pd.DataFrame, start_row=1, start_col=1, max_rows=None):
    data = df.head(max_rows) if max_rows else df
    for j, col in enumerate(data.columns, start_col):
        c = ws.cell(start_row, j, col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, row) in enumerate(data.iterrows(), start_row + 1):
        for j, col in enumerate(data.columns, start_col):
            v = row[col]
            if pd.isna(v):
                v = ""
            ws.cell(i, j, v)
    return start_row + len(data) + 1


def style_sheet(ws, freeze="A2"):
    ws.freeze_panes = freeze
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col[:100]:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 40)
    ws.auto_filter.ref = ws.dimensions


def kpi(ws, row: int, col: int, name: str, value: str, fill="D9EAD3"):
    ws.cell(row, col, name)
    ws.cell(row + 1, col, value)
    for r in [row, row + 1]:
        ws.cell(r, col).fill = PatternFill("solid", fgColor=fill)
        ws.cell(r, col).alignment = Alignment(horizontal="center")
    ws.cell(row, col).font = Font(bold=True, color="404040")
    ws.cell(row + 1, col).font = Font(bold=True, size=15, color="203864")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    backend = read_backend()
    promoted = read_ad(PROMOTED_PATH)
    search = read_ad(SEARCH_PATH)

    backend["结构角色"] = backend.apply(account_role, axis=1)
    actions = backend.apply(budget_action, axis=1)
    backend["预算动作"] = [a[0] for a in actions]
    backend["广告策略"] = [a[1] for a in actions]
    backend["经营目标"] = [a[2] for a in actions]
    backend["销售占比"] = backend["销售额"] / backend["销售额"].sum()
    backend["广告花费占比"] = backend["广告花费"] / backend["广告花费"].sum()
    backend["利润贡献占比"] = backend["毛利润"] / backend["毛利润"].sum()

    parent = (
        backend.groupby("父ASIN", dropna=False)
        .agg(
            ASIN数=("ASIN", "count"),
            销售额=("销售额", "sum"),
            销量=("销量", "sum"),
            订单量=("订单量", "sum"),
            广告花费=("广告花费", "sum"),
            毛利润=("毛利润", "sum"),
            可售=("可售", "sum"),
            自然点击量=("自然点击量", "sum"),
            广告曝光量=("广告曝光量", "sum"),
            广告点击量=("广告点击量", "sum"),
            访问次数=("访问次数", "sum"),
        )
        .reset_index()
    )
    parent["毛利率"] = parent["毛利润"] / parent["销售额"].replace(0, pd.NA)
    parent["ACoAS"] = parent["广告花费"] / parent["销售额"].replace(0, pd.NA)
    parent["销售占比"] = parent["销售额"] / parent["销售额"].sum()
    parent["广告花费占比"] = parent["广告花费"] / parent["广告花费"].sum()
    parent["利润贡献占比"] = parent["毛利润"] / parent["毛利润"].sum()
    parent = parent.sort_values("销售额", ascending=False)

    same_parent = backend[backend["父ASIN"].astype(str) == TARGET_PARENT].copy()
    same_parent = same_parent.sort_values("销售额", ascending=False)

    ad_asin = (
        promoted.groupby(["广告ASIN", "广告SKU"], dropna=False)
        .agg(展示量=("展示量", "sum"), 点击量=("点击量", "sum"), 花费=("花费", "sum"), 广告销售额=("7天总销售额", "sum"), 广告订单=("7天总订单数(#)", "sum"))
        .reset_index()
    )
    ad_asin["ACOS"] = ad_asin["花费"] / ad_asin["广告销售额"].replace(0, pd.NA)
    ad_asin["CVR"] = ad_asin["广告订单"] / ad_asin["点击量"].replace(0, pd.NA)
    ad_asin = ad_asin.sort_values("花费", ascending=False)

    search_campaign = (
        search.groupby(["广告活动名称"], dropna=False)
        .agg(展示量=("展示量", "sum"), 点击量=("点击量", "sum"), 花费=("花费", "sum"), 广告销售额=("7天总销售额", "sum"), 广告订单=("7天总订单数(#)", "sum"))
        .reset_index()
    )
    search_campaign["ACOS"] = search_campaign["花费"] / search_campaign["广告销售额"].replace(0, pd.NA)
    search_campaign["CVR"] = search_campaign["广告订单"] / search_campaign["点击量"].replace(0, pd.NA)
    search_campaign["结构判断"] = search_campaign.apply(
        lambda r: "高花费低回报，收缩" if r["花费"] > 100 and (pd.isna(r["ACOS"]) or r["ACOS"] > 0.4) else ("可保留/优化" if r["广告订单"] > 0 else "低量观察"),
        axis=1,
    )
    search_campaign = search_campaign.sort_values("花费", ascending=False)

    issues = pd.DataFrame(
        [
            ["父体预算错配", "高", "B0DHJK7NNF广告花费最高且毛利润为负，同父体利润被吞噬", "将其预算下调50%-80%，只保留能证明盈利的精准词"],
            ["库存错配", "高", "B0F2LQNTDF、B0GTZ1CGTD、B0F2LLRF87均低于14天或断货", "低库存款控量，高库存款用低成本促销/长尾词承接"],
            ["高库存低动销", "中高", "多个同父体变体可售天数超过180天但销售额低", "避免高CPC强推；先做价格、coupon、图片和变体排序优化"],
            ["广告依赖过高", "中高", "B0DHJK7NNF、B0GTZ1CGTD广告订单占比超过100%，归因和自然承接异常", "降低泛流量预算，重建自然流量和精准词结构"],
            ["数据结构异常", "中", "B0D943FGNY在后台出现不同父ASIN记录", "先核对父体/变体映射，避免广告和库存判断错配"],
            ["同系列自然流量分散", "中", "父体B0FXRZ3YMT有17个ASIN，部分自然点击高但转化低", "优化变体排序，把高转化/高毛利颜色前置"],
        ],
        columns=["结构问题", "优先级", "证据", "优化方案"],
    )

    budget = backend[
        [
            "ASIN",
            "父ASIN",
            "结构角色",
            "销售额",
            "广告花费",
            "毛利润",
            "毛利率",
            "ACoAS",
            "ACoS",
            "可售天数",
            "自然点击量",
            "访问转化率",
            "预算动作",
            "广告策略",
            "经营目标",
        ]
    ].sort_values(["结构角色", "销售额"], ascending=[True, False])

    stock = backend[
        [
            "ASIN",
            "父ASIN",
            "销售额",
            "30日销量",
            "日均销量",
            "可售",
            "可售天数",
            "总可售天数",
            "毛利率",
            "ACoAS",
            "结构角色",
            "预算动作",
        ]
    ].sort_values("可售天数")

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = "全账户 / 同父体 / 同系列结构问题分析"
    ws["A1"].font = Font(bold=True, size=18, color="1F4E78")
    ws.merge_cells("A1:H1")
    ws["A2"] = "数据来源：后台商品分析V2 2026-05-07至2026-06-05；商品推广广告报告约30天。广告与后台为不同归因口径，不强行对齐。"
    ws.merge_cells("A2:H2")
    ws["A2"].alignment = Alignment(wrap_text=True)

    kpi(ws, 4, 1, "全账户销售额", money(backend["销售额"].sum()))
    kpi(ws, 4, 2, "全账户毛利润", money(backend["毛利润"].sum()), "E2F0D9")
    kpi(ws, 4, 3, "全账户毛利率", pct(backend["毛利润"].sum() / backend["销售额"].sum()), "FFF2CC")
    kpi(ws, 4, 4, "后台广告花费", money(backend["广告花费"].sum()), "DDEBF7")
    kpi(ws, 4, 5, "全账户ACoAS", pct(backend["广告花费"].sum() / backend["销售额"].sum()), "FCE4D6")
    kpi(ws, 4, 6, "ASIN数", str(len(backend)), "D9EAD3")
    kpi(ws, 4, 7, "主父体销售占比", pct(parent.iloc[0]["销售占比"]), "E2F0D9")
    kpi(ws, 4, 8, "止损ASIN数", str((backend["结构角色"] == "止损收缩").sum()), "F4CCCC")
    write_df(ws, issues, 8, 1)

    chart_source = parent[["父ASIN", "销售额", "广告花费", "毛利润"]].head(5)
    start = 17
    write_df(ws, chart_source, start, 1)
    bar = BarChart()
    bar.title = "父体销售/广告/利润结构"
    bar.y_axis.title = "USD"
    data = Reference(ws, min_col=2, max_col=4, min_row=start, max_row=start + len(chart_source))
    cats = Reference(ws, min_col=1, min_row=start + 1, max_row=start + len(chart_source))
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.height = 7
    bar.width = 15
    ws.add_chart(bar, "F17")

    role_counts = backend.groupby("结构角色").agg(ASIN数=("ASIN", "count")).reset_index()
    role_start = 25
    write_df(ws, role_counts, role_start, 1)
    doughnut = DoughnutChart()
    doughnut.title = "ASIN结构角色分布"
    d_data = Reference(ws, min_col=2, min_row=role_start, max_row=role_start + len(role_counts))
    d_cats = Reference(ws, min_col=1, min_row=role_start + 1, max_row=role_start + len(role_counts))
    doughnut.add_data(d_data, titles_from_data=True)
    doughnut.set_categories(d_cats)
    doughnut.height = 7
    doughnut.width = 10
    ws.add_chart(doughnut, "F26")

    for name, df in [
        ("结构问题矩阵", issues),
        ("父体结构", parent),
        ("同父体变体策略", same_parent),
        ("全账户ASIN策略", budget),
        ("预算重分配建议", budget),
        ("库存风险&清库存", stock),
        ("广告ASIN归因", ad_asin),
        ("活动结构诊断", search_campaign),
        ("后台原始数据", backend),
    ]:
        w = wb.create_sheet(name)
        write_df(w, df)
        style_sheet(w)

    style_sheet(ws, "A8")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["C"].width = 44
    ws.column_dimensions["D"].width = 48
    ws.row_dimensions[2].height = 34

    output_xlsx = OUTPUT_DIR / "全账户同父体同系列结构问题分析.xlsx"
    wb.save(output_xlsx)

    parent_main = parent.iloc[0]
    b0dh = backend[backend["ASIN"].astype(str) == "B0DHJK7NNF"].iloc[0]
    b0f2 = backend[backend["ASIN"].astype(str) == TARGET_ASIN].iloc[0]
    gt = backend[backend["ASIN"].astype(str) == "B0GTZ1CGTD"].iloc[0]
    overstock_count = int(((backend["可售天数"].fillna(0) > 180) & (backend["销售额"].fillna(0) < 500)).sum())
    low_stock_count = int((backend["可售天数"].fillna(999) < 14).sum())

    report = f"""# 全账户 / 同父体 / 同系列结构问题分析与优化方案

## 1. 总体判断
全账户近30天后台销售额约 {money(backend['销售额'].sum())}，毛利润 {money(backend['毛利润'].sum())}，整体毛利率 {pct(backend['毛利润'].sum() / backend['销售额'].sum())}，后台广告花费 {money(backend['广告花费'].sum())}，ACoAS {pct(backend['广告花费'].sum() / backend['销售额'].sum())}。账户不是没有利润，但利润结构很不健康：少数SKU在赚钱，少数高花费SKU在吞利润，库存又同时存在“断货款”和“积压款”。

## 2. 全账户结构问题
- 广告预算集中在低利润/亏损款：B0DHJK7NNF 花费 {money(b0dh['广告花费'])}，毛利润 {money(b0dh['毛利润'])}，ACoAS {pct(b0dh['ACoAS'])}，应作为第一止损对象。
- 高销量款库存不足：B0F2LQNTDF 销售额 {money(b0f2['销售额'])}，毛利润 {money(b0f2['毛利润'])}，但可售仅 {float(b0f2['可售天数']):.0f} 天；B0GTZ1CGTD 可售仅 {float(gt['可售天数']):.0f} 天。
- 库存两极化：低于14天库存风险的ASIN有 {low_stock_count} 个，同时可售天数>180且销售额<500的积压ASIN有 {overstock_count} 个。
- 广告依赖异常：B0DHJK7NNF、B0GTZ1CGTD广告订单占比超过100%，说明自然流量承接弱，或广告归因集中在少数变体上。
- 数据映射需要清理：B0D943FGNY出现不同父ASIN记录，建议先核对变体关系和后台数据导出口径。

## 3. 同父体 B0FXRZ3YMT 问题
主父体 {parent_main['父ASIN']} 包含 {int(parent_main['ASIN数'])} 个ASIN，贡献销售额 {money(parent_main['销售额'])}，占全账户 {pct(parent_main['销售占比'])}，广告花费 {money(parent_main['广告花费'])}，毛利润 {money(parent_main['毛利润'])}。问题不是父体没有市场，而是父体内部预算和库存错配。

优先级：
1. B0DHJK7NNF：立即收缩50%-80%广告预算，只保留低ACOS精准词。
2. B0F2LQNTDF：利润正、自然点击强，但库存低，先控量防断货，库存恢复到21天后再放量。
3. B0F2MGP1WN：毛利率 {pct(backend.loc[backend['ASIN'].eq('B0F2MGP1WN'), '毛利率'].iloc[0])}、ACoAS {pct(backend.loc[backend['ASIN'].eq('B0F2MGP1WN'), 'ACoAS'].iloc[0])}，库存43天，是更适合承接增量的利润型变体。
4. B0F2MBQVNN、B0F2MGST2F、B0F2M53ZX4：毛利率高但广告弱，可用低CPC精准长尾词测试，不适合大预算泛投。
5. B0F2M83HQJ、B0D943FGNY：高库存、低转化、广告效率差，先修Listing和价格，暂停高价广告。

## 4. 预算重分配方案
- 削减池：B0DHJK7NNF、B0GTZ1CGTD、B0F2M83HQJ、B0D943FGNY，削减比例50%-80%，释放低效广告费。
- 保守池：B0F2LQNTDF、B0F2LLRF87，因库存不足，只保留高转化精准词和品牌防守词。
- 放量池：B0F2MGP1WN、B0F2MGST2F、B0F2MBQVNN、B0F2M53ZX4，前提是主图/变体顺序和coupon到位，采用低CPC精准/长尾词测试。
- 清库存池：可售天数>180的低动销变体，不做高CPC抢词，优先用coupon、价格梯度、变体排序、捆绑场景和低价长尾词处理。

## 5. 30天执行节奏
- 第1-3天：砍B0DHJK7NNF等亏损款预算；核对B0D943FGNY父体映射；给低库存款设置预算上限。
- 第4-7天：把B0F2LQNTDF只保留盈利词；把预算迁移到B0F2MGP1WN等利润型变体；检查库存补货节点。
- 第2周：调整父体变体排序，让高利润/库存健康款获得更多自然曝光；对积压款做coupon和主图测试。
- 第3-4周：库存恢复后，B0F2LQNTDF和利润型变体逐步+20%-40%预算；每7天复盘ACoAS、毛利率、库存天数。

详细表格见工作簿：结构问题矩阵、父体结构、同父体变体策略、全账户ASIN策略、预算重分配建议、库存风险&清库存、广告ASIN归因、活动结构诊断。
"""
    output_md = OUTPUT_DIR / "全账户同父体同系列结构问题分析.md"
    output_md.write_text(report, encoding="utf-8-sig")

    print(f"XLSX={output_xlsx}")
    print(f"MD={output_md}")
    print(f"backend_rows={len(backend)}")
    print(f"parents={len(parent)}")
    print(f"same_parent_rows={len(same_parent)}")
    print(f"low_stock={low_stock_count}, overstock={overstock_count}")


if __name__ == "__main__":
    main()
