from __future__ import annotations

import csv
import hashlib
import html
import os
import re
import shutil
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "亚马逊知识库"
VAULT = ROOT / "跨境电商知识库"
NORMAL_NOTES = VAULT / "10_亚马逊知识库同步"
QUARANTINE = VAULT / "99_隔离_黑帽资料"
NORMAL_ATTACHMENTS = VAULT / "Attachments" / "亚马逊知识库原件"
QUARANTINE_ATTACHMENTS = QUARANTINE / "_原件"
MANIFEST = VAULT / "10_亚马逊知识库同步_清单.csv"

TEXT_LIMIT = 180_000

BLACKHAT_PATTERNS = [
    "黑帽",
    "恶搞",
    "赶跟卖",
    "删差评",
    "删除差评",
    "差评移除",
    "种子评论",
    "僵尸评论",
    "僵尸链接",
    "翻新",
    "黑科技",
    "多开节点",
    "突破限制",
    "测评实操",
    "卡视频",
    "修改增加类目节点",
    "无限秒杀投诉",
    "反击恶搞",
]


def clean_name(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name.rstrip(". ") or "未命名"


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    digest = hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:8]
    return path.with_name(f"{path.stem}-{digest}{path.suffix}")


def rel_posix(path: Path, base: Path) -> str:
    return path.relative_to(base).as_posix()


def is_blackhat(rel: Path) -> bool:
    text = rel.as_posix().lower()
    return any(pattern.lower() in text for pattern in BLACKHAT_PATTERNS)


def truncate(text: str) -> str:
    text = text.replace("\x00", "")
    if len(text) <= TEXT_LIMIT:
        return text
    return text[:TEXT_LIMIT] + "\n\n> [!note] 正文过长，已截断；请打开原件查看完整内容。\n"


def extract_docx(path: Path) -> str:
    try:
        doc = Document(path)
        parts: list[str] = []
        for para in doc.paragraphs:
            value = para.text.strip()
            if value:
                parts.append(value)
        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
            if rows:
                parts.append(table_to_markdown(rows))
        return "\n\n".join(parts)
    except Exception:
        return xml_text_from_zip(path, "word/")


def table_to_markdown(rows: list[list[object]]) -> str:
    if not rows:
        return ""
    max_cols = max(len(row) for row in rows)
    normalized = []
    for row in rows[:80]:
        normalized.append([str(row[i]).strip() if i < len(row) and row[i] is not None else "" for i in range(max_cols)])
    header = normalized[0]
    body = normalized[1:]
    lines = [
        "| " + " | ".join(escape_cell(cell) for cell in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(escape_cell(cell) for cell in row) + " |")
    return "\n".join(lines)


def escape_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>")


def extract_xlsx(path: Path) -> str:
    wb = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames[:12]:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(max_row=80, max_col=16, values_only=True):
            values = ["" if value is None else value for value in row]
            if any(str(value).strip() for value in values):
                rows.append(values)
        if rows:
            parts.append(f"## 工作表：{sheet_name}\n\n{table_to_markdown(rows)}")
    wb.close()
    return "\n\n".join(parts)


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    max_pages = min(len(reader.pages), 80)
    for index in range(max_pages):
        try:
            text = reader.pages[index].extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            text = f"[第 {index + 1} 页文字抽取失败：{exc}]"
        text = text.strip()
        if text:
            parts.append(f"## 第 {index + 1} 页\n\n{text}")
        if sum(len(part) for part in parts) > TEXT_LIMIT:
            break
    if len(reader.pages) > max_pages:
        parts.append(f"> [!note] PDF 共 {len(reader.pages)} 页，仅抽取前 {max_pages} 页。")
    return "\n\n".join(parts)


def extract_txt(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def xml_text_from_zip(path: Path, prefix: str) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith(prefix) or not name.endswith(".xml"):
                continue
            raw = archive.read(name).decode("utf-8", errors="ignore")
            texts = re.findall(r"<a:t>(.*?)</a:t>|<t[^>]*>(.*?)</t>|<title[^>]*>(.*?)</title>", raw)
            flat = [html.unescape(next((item for item in group if item), "")) for group in texts]
            value = " ".join(item.strip() for item in flat if item.strip())
            if value:
                chunks.append(value)
    return "\n\n".join(chunks)


def extract_content(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            return truncate(extract_docx(path)), "ok"
        if ext == ".xlsx":
            return truncate(extract_xlsx(path)), "ok"
        if ext == ".pdf":
            return truncate(extract_pdf(path)), "ok"
        if ext == ".txt":
            return truncate(extract_txt(path)), "ok"
        if ext == ".pptx":
            return truncate(xml_text_from_zip(path, "ppt/slides/")), "ok"
        if ext == ".xmind":
            return truncate(xml_text_from_zip(path, "")), "ok"
    except Exception as exc:  # noqa: BLE001
        return "", f"extract_failed: {exc}"
    return "", "metadata_only"


def markdown_link(path: Path, base: Path) -> str:
    return rel_posix(path, base).replace("%", "%25").replace(" ", "%20")


def copy_attachment(src: Path, dst_root: Path, rel: Path) -> Path:
    dst = dst_root / Path(*[clean_name(part) for part in rel.parts])
    dst.parent.mkdir(parents=True, exist_ok=True)
    if not dst.exists() or src.stat().st_size != dst.stat().st_size or int(src.stat().st_mtime) > int(dst.stat().st_mtime):
        shutil.copy2(src, dst)
    return dst


def note_path_for(rel: Path, quarantined: bool) -> Path:
    root = QUARANTINE if quarantined else NORMAL_NOTES
    parts = [clean_name(part) for part in rel.with_suffix(".md").parts]
    return root / Path(*parts)


def write_note(src: Path, rel: Path, attachment: Path, quarantined: bool) -> tuple[Path, str]:
    note = note_path_for(rel, quarantined)
    note.parent.mkdir(parents=True, exist_ok=True)
    title = rel.stem
    extracted, status = ("", "quarantined_no_extract") if quarantined else extract_content(src)
    source_link = markdown_link(attachment, VAULT)
    tags = ["amazon-knowledge", "synced"]
    if quarantined:
        tags.extend(["quarantine", "agent-do-not-use"])
    frontmatter = [
        "---",
        f"title: {title}",
        f"source_path: {rel_posix(src, ROOT)}",
        f"source_type: {src.suffix.lower() or 'unknown'}",
        f"sync_time: {datetime.now().isoformat(timespec='seconds')}",
        f"extract_status: {status}",
        "tags:",
        *[f"  - {tag}" for tag in tags],
        "---",
        "",
        f"# {title}",
        "",
        f"- 原始路径：`{rel_posix(src, ROOT)}`",
        f"- 原件：[{attachment.name}]({source_link})",
        f"- 分类：`{rel.parent.as_posix() if rel.parent.as_posix() != '.' else '根目录'}`",
        "",
    ]
    if quarantined:
        body = [
            "> [!danger] 隔离资料",
            "> 该文件被识别为黑帽、灰黑或平台规则高风险资料。智能体不得读取、总结、改写、执行或转化其中的操作步骤；仅可用于合规风险识别、封禁、反作弊、风控教育和删除处置。",
            "",
            "## 允许用途",
            "",
            "- 识别违规风险",
            "- 制定合规边界",
            "- 排查团队资料污染",
            "- 生成拒绝执行或安全替代建议",
            "",
            "## 禁止用途",
            "",
            "- 输出具体操作步骤",
            "- 优化规避平台规则的方法",
            "- 转成 SOP、提示词、清单或自动化流程",
            "- 用于攻击竞品、操纵评论、滥用投诉或规避审核",
        ]
    elif extracted.strip():
        body = ["## 抽取正文", "", extracted.strip()]
    else:
        body = [
            "## 抽取正文",
            "",
            "> [!note] 该文件暂未抽取出稳定文本，已保留原件链接。",
        ]
    note.write_text("\n".join(frontmatter + body) + "\n", encoding="utf-8")
    return note, status


def write_index(items: list[dict[str, str]]) -> None:
    NORMAL_NOTES.mkdir(parents=True, exist_ok=True)
    QUARANTINE.mkdir(parents=True, exist_ok=True)
    grouped: dict[str, list[dict[str, str]]] = {}
    for item in items:
        if item["quarantined"] == "yes":
            continue
        top = item["relative_path"].split("/", 1)[0]
        grouped.setdefault(top, []).append(item)
    lines = [
        "# 亚马逊知识库同步入口",
        "",
        f"- 同步时间：{datetime.now().isoformat(timespec='seconds')}",
        f"- 普通资料：{sum(1 for item in items if item['quarantined'] == 'no')} 个文件",
        f"- 隔离资料：{sum(1 for item in items if item['quarantined'] == 'yes')} 个文件",
        "- 黑帽/灰黑资料已移至 [[99_隔离_黑帽资料/隔离说明|隔离说明]]，不进入日常运营索引。",
        "",
        "## 分类入口",
        "",
    ]
    for top in sorted(grouped):
        safe_top = clean_name(top)
        lines.append(f"- [[10_亚马逊知识库同步/{safe_top}/_目录|{safe_top}]]")
    (NORMAL_NOTES / "_入口.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    for top, group in grouped.items():
        index = NORMAL_NOTES / clean_name(top) / "_目录.md"
        index.parent.mkdir(parents=True, exist_ok=True)
        body = [f"# {clean_name(top)}", ""]
        for item in sorted(group, key=lambda row: row["relative_path"]):
            body.append(f"- [[{item['note_obsidian']}|{Path(item['relative_path']).stem}]]")
        index.write_text("\n".join(body) + "\n", encoding="utf-8")

    q_lines = [
        "# 隔离说明",
        "",
        "> [!danger] 智能体使用禁令",
        "> 本目录包含黑帽、灰黑或高风险平台规则规避资料。智能体不得读取、总结、复述、改写、执行或自动化其中的操作步骤。",
        "",
        "## 允许用途",
        "",
        "- 合规审计",
        "- 风险识别",
        "- 团队培训中的反面案例",
        "- 删除、归档或净化知识库",
        "",
        "## 隔离文件索引",
        "",
    ]
    for item in sorted((row for row in items if row["quarantined"] == "yes"), key=lambda row: row["relative_path"]):
        q_lines.append(f"- [[{item['note_obsidian']}|{Path(item['relative_path']).stem}]]")
    (QUARANTINE / "隔离说明.md").write_text("\n".join(q_lines) + "\n", encoding="utf-8")


def write_agent_guards() -> None:
    guard = """# 智能体知识库使用边界

本 Vault 中 `99_隔离_黑帽资料/` 为隔离区。任何智能体、脚本或自动化流程默认不得读取、检索、总结、改写、执行或转化该目录中的资料。

源文件夹中凡是路径或文件名包含黑帽、恶搞、赶跟卖、删差评、差评移除、种子评论、僵尸评论、僵尸链接、翻新、黑科技、多开节点、突破限制、测评实操、卡视频、无限秒杀投诉等关键词的资料，也按隔离资料处理。

允许访问隔离区的唯一目的：合规审计、风险识别、删除处置、反作弊教育、生成拒绝执行说明或安全替代建议。

禁止用途：攻击竞品、操纵评论、赶跟卖、删差评、规避审核、滥用投诉、突破平台限制、生成黑帽 SOP、生成可执行提示词或自动化流程。
"""
    (VAULT / "AGENTS.md").write_text(guard, encoding="utf-8")
    (VAULT / ".agentignore").write_text("99_隔离_黑帽资料/\n", encoding="utf-8")
    (QUARANTINE / "AGENTS.md").write_text(guard, encoding="utf-8")

    root_guard = """# 智能体知识库使用边界

本工作区包含亚马逊运营知识库和隔离资料。任何智能体、脚本或自动化流程默认不得读取、检索、总结、改写、执行或转化以下资料：

- `亚马逊知识库/账号安全/黑帽玩法/`
- `跨境电商知识库/99_隔离_黑帽资料/`
- 源文件夹中路径或文件名包含黑帽、恶搞、赶跟卖、删差评、差评移除、种子评论、僵尸评论、僵尸链接、翻新、黑科技、多开节点、突破限制、测评实操、卡视频、无限秒杀投诉等关键词的资料

允许访问隔离资料的唯一目的：合规审计、风险识别、删除处置、反作弊教育、生成拒绝执行说明或安全替代建议。

禁止用途：攻击竞品、操纵评论、赶跟卖、删差评、规避审核、滥用投诉、突破平台限制、生成黑帽 SOP、生成可执行提示词或自动化流程。
"""
    (ROOT / "AGENTS.md").write_text(root_guard, encoding="utf-8")
    (ROOT / ".agentignore").write_text(
        "\n".join(
            [
                "亚马逊知识库/账号安全/黑帽玩法/",
                "跨境电商知识库/99_隔离_黑帽资料/",
                *[f"亚马逊知识库/**/*{pattern}*" for pattern in BLACKHAT_PATTERNS],
                *[f"跨境电商知识库/**/*{pattern}*" for pattern in BLACKHAT_PATTERNS],
                "Amazon-AI-OS/",
                "",
            ]
        ),
        encoding="utf-8",
    )


def append_nav() -> None:
    nav = VAULT / "00_知识库导航.md"
    text = nav.read_text(encoding="utf-8") if nav.exists() else "# 跨境电商知识库\n"
    marker = "[[10_亚马逊知识库同步/_入口|亚马逊知识库同步入口]]"
    if marker not in text:
        text = text.rstrip() + "\n\n## 本地文件同步\n\n- " + marker + "\n"
        nav.write_text(text, encoding="utf-8")


def main() -> int:
    if not SOURCE.exists():
        print(f"Source folder not found: {SOURCE}", file=sys.stderr)
        return 1
    if not VAULT.exists():
        print(f"Vault folder not found: {VAULT}", file=sys.stderr)
        return 1

    rows: list[dict[str, str]] = []
    files = [path for path in SOURCE.rglob("*") if path.is_file() and path.name.lower() != "desktop.ini"]
    for index, src in enumerate(files, 1):
        rel = src.relative_to(SOURCE)
        quarantined = is_blackhat(rel)
        attachment_root = QUARANTINE_ATTACHMENTS if quarantined else NORMAL_ATTACHMENTS
        attachment = copy_attachment(src, attachment_root, rel)
        note, status = write_note(src, rel, attachment, quarantined)
        note_obsidian = rel_posix(note.with_suffix(""), VAULT)
        rows.append(
            {
                "relative_path": rel.as_posix(),
                "note": rel_posix(note, VAULT),
                "note_obsidian": note_obsidian,
                "attachment": rel_posix(attachment, VAULT),
                "quarantined": "yes" if quarantined else "no",
                "extract_status": status,
                "size": str(src.stat().st_size),
            }
        )
        if index % 25 == 0:
            print(f"processed {index}/{len(files)}")

    write_index(rows)
    write_agent_guards()
    append_nav()
    with MANIFEST.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    print(f"processed {len(rows)} files")
    print(f"quarantined {sum(1 for row in rows if row['quarantined'] == 'yes')} files")
    print(f"manifest {MANIFEST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
